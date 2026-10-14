"""Food library → Import from Excel (2 steps: check, then import).

Columns are matched by name. Required: Food, Portion_g, Calories_kcal. Everything is stored per 100 g.
"""
import csv
import io
import re
from decimal import Decimal, InvalidOperation

from django.db import transaction

from .models import FOOD_GROUPS, Diet, Food, FoodImportJob, FoodUnit

MAX_ROWS = 5000
REQUIRED = ["food", "portion_g", "calories_kcal"]
ALIASES = {
    "food": ["food", "name", "food_name"],
    "portion_g": ["portion_g", "grams", "weight_g", "serving_g"],
    "calories_kcal": ["calories_kcal", "kcal", "calories", "energy_kcal"],
    "portion": ["portion", "quantity", "portion/quantity", "unit"],
    "protein": ["protein", "protein_g"],
    "carbs": ["carbs", "carbs_g", "carbohydrate", "carbohydrates_g"],
    "fat": ["fat", "fat_g", "total_fat"],
    "fibre": ["fibre", "fiber", "fibre_g", "fiber_g"],
    "sodium_mg": ["sodium", "sodium_mg"],
    "iron_mg": ["iron", "iron_mg"],
    "calcium_mg": ["calcium", "calcium_mg"],
    "potassium_mg": ["potassium", "potassium_mg"],
    "group": ["group", "category", "food_group"],
    "diet": ["diet", "type", "veg/non-veg"],
}
GROUP_MAP = {"grains": "carbs", "cereals": "carbs", "egg": "protein", "eggs": "protein", "meat": "protein",
             "pulses": "protein", "non-veg": "indian", "dishes": "indian", "beverages": "beverages",
             "vegetables": "veg", "fruit": "fruits", "milk": "dairy", "nuts": "fats", "oils": "fats"}
DIET_MAP = {"vegetarian": "veg", "veg": "veg", "non-vegetarian": "nonveg", "non-veg": "nonveg", "nonveg": "nonveg",
            "dairy": "veg", "egg": "egg", "eggetarian": "egg", "vegan": "vegan"}
UNIT_RE = re.compile(r"^(?P<name>.+?)\s*\((?P<g>[\d.]+)\s*g\)\s*$")


def _norm(h):
    return re.sub(r"\s+", "_", str(h or "").strip().lower())


def _num(v):
    if v in (None, ""):
        return None
    try:
        return Decimal(str(v).replace(",", "").strip())
    except InvalidOperation:
        return "bad"


def read_rows(f):
    name = f.name.lower()
    if name.endswith(".csv"):
        text = f.read().decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(text)))
    elif name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(f, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        raise ValueError("Upload an Excel (.xlsx) or CSV file.")
    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        raise ValueError("The file is empty.")
    if len(rows) - 1 > MAX_ROWS:
        raise ValueError(f"Up to {MAX_ROWS} rows at a time. Split the file.")
    header = [_norm(h) for h in rows[0]]
    index = {}
    for key, names in ALIASES.items():
        for i, h in enumerate(header):
            if h in names:
                index[key] = i
                break
    missing = [k for k in REQUIRED if k not in index]
    if missing:
        raise ValueError("Missing columns: " + ", ".join(m.replace("_", " ").title() for m in missing))
    return index, rows[1:]


def check(f, *, user):
    index, raw = read_rows(f)
    existing = {n.lower(): fid for fid, n in Food.objects.filter(owner_client__isnull=True)
                .values_list("id", "name")}
    out, seen = [], set()
    egg_rows = 0
    for n, r in enumerate(raw, start=2):
        def cell(k):
            i = index.get(k)
            return r[i] if i is not None and i < len(r) else None

        name = " ".join(str(cell("food") or "").split())[:120]
        portion = _num(cell("portion_g"))
        kcal = _num(cell("calories_kcal"))
        errors = []
        if not name:
            errors.append("No food name")
        if portion in (None, "bad") or (portion is not None and portion != "bad" and portion <= 0):
            errors.append("Portion_g must be a number above 0")
        if kcal in (None, "bad"):
            errors.append("Calories_kcal must be a number")
        macros = {k: _num(cell(k)) for k in ("protein", "carbs", "fat", "fibre", "sodium_mg", "iron_mg",
                                              "calcium_mg", "potassium_mg")}
        for k, v in macros.items():
            if v == "bad":
                errors.append(f"{k} is not a number")
        row = {"line": n, "name": name, "errors": errors}
        if not errors:
            f100 = Decimal(100) / portion
            per100 = {k: float(round(v * f100, 2)) for k, v in macros.items() if v not in (None, "bad")}
            per100["kcal"] = float(round(kcal * f100, 1))
            calc = 4 * per100.get("protein", 0) + 4 * per100.get("carbs", 0) + 9 * per100.get("fat", 0)
            diet_raw = str(cell("diet") or "").strip().lower()
            group_raw = str(cell("group") or "").strip().lower()
            is_egg = "egg" in name.lower()
            if is_egg and DIET_MAP.get(diet_raw) == "nonveg":
                egg_rows += 1
            unit = None
            m = UNIT_RE.match(str(cell("portion") or ""))
            if m:
                unit = {"name": m.group("name").strip()[:40], "grams": float(m.group("g"))}
            elif cell("portion"):
                unit = {"name": str(cell("portion")).strip()[:40], "grams": float(portion)}
            key = name.lower()
            row.update({
                "per100": per100, "calc_kcal": round(calc), "kcal_differs": abs(calc - per100["kcal"]) > max(
                    15, 0.1 * per100["kcal"]) and any(k in per100 for k in ("protein", "carbs", "fat")),
                "diet_raw": diet_raw, "group_raw": group_raw, "is_egg": is_egg, "unit": unit,
                "existing": str(existing[key]) if key in existing else None,
                "dupe_in_file": key in seen,
            })
            seen.add(key)
        out.append(row)
    ok = [r for r in out if not r["errors"]]
    summary = {
        "rows": len(out), "new": sum(1 for r in ok if not r["existing"]),
        "already_there": sum(1 for r in ok if r["existing"]),
        "kcal_differs": sum(1 for r in ok if r["kcal_differs"]),
        "errors": sum(1 for r in out if r["errors"]), "egg_rows_nonveg": egg_rows,
        "groups_found": sorted({r["group_raw"] for r in ok if r["group_raw"]}),
        "diets_found": sorted({r["diet_raw"] for r in ok if r["diet_raw"]}),
        "group_suggestions": {g: GROUP_MAP.get(g, g if g in FOOD_GROUPS else "other")
                              for g in {r["group_raw"] for r in ok if r["group_raw"]}},
        "diet_suggestions": {d: DIET_MAP.get(d, "veg") for d in {r["diet_raw"] for r in ok if r["diet_raw"]}},
        "units": sum(1 for r in ok if r.get("unit")),
    }
    f.seek(0)
    job = FoodImportJob.objects.create(file=f, rows=out, summary=summary, created_by=user)
    return job


@transaction.atomic
def run(job, *, duplicates="skip", kcal_choice=None, eggs_eggetarian=True, group_map=None, diet_map=None,
        user=None):
    """duplicates: skip | update | new. kcal_choice: {line: "sheet" | "macros"}."""
    if job.status != FoodImportJob.Status.CHECKED:
        raise ValueError("This file was imported already.")
    kcal_choice = kcal_choice or {}
    group_map = {**job.summary.get("group_suggestions", {}), **(group_map or {})}
    diet_map = {**job.summary.get("diet_suggestions", {}), **(diet_map or {})}
    created, updated, skipped = [], [], 0
    previous = {}
    for r in job.rows:
        if r["errors"] or r.get("dupe_in_file"):
            skipped += 1
            continue
        p = r["per100"]
        kcal = p["kcal"]
        if r["kcal_differs"] and kcal_choice.get(str(r["line"])) == "macros":
            kcal = r["calc_kcal"]
        diet = diet_map.get(r["diet_raw"], "veg") if r["diet_raw"] else "veg"
        if r["is_egg"] and eggs_eggetarian and diet == "nonveg":
            diet = Diet.EGG
        group = group_map.get(r["group_raw"], "other") if r["group_raw"] else "other"
        if group not in FOOD_GROUPS:
            group = "other"
        fields = dict(kcal=min(Decimal(str(kcal)), Decimal(900)), diet=diet, group=group,
                      **{k: Decimal(str(p[k])) for k in ("protein", "carbs", "fat", "fibre", "sodium_mg", "iron_mg",
                                                         "calcium_mg", "potassium_mg") if k in p})
        if r["existing"] and duplicates == "skip":
            skipped += 1
            continue
        if r["existing"] and duplicates == "update":
            food = Food.objects.get(id=r["existing"])
            previous[str(food.id)] = {k: str(getattr(food, k)) if getattr(food, k) is not None else None
                                      for k in fields}
            for k, v in fields.items():
                setattr(food, k, v)
            food.save()
            updated.append(food.id)
        else:
            food = Food.objects.create(name=r["name"], source=Food.Source.IMPORTED, source_note="Excel",
                                       created_by=user, **fields)
            created.append(food.id)
        if r.get("unit"):
            unit, _ = FoodUnit.objects.get_or_create(food=food, name=r["unit"]["name"],
                                                     defaults={"grams": Decimal(str(r["unit"]["grams"]))})
            if food.default_unit_id is None:
                food.default_unit = unit
                food.save(update_fields=["default_unit"])
    job.status = FoodImportJob.Status.IMPORTED
    job.summary = {**job.summary, "result": {"created": [str(i) for i in created],
                                             "updated": [str(i) for i in updated], "skipped": skipped,
                                             "previous": previous}}
    job.save(update_fields=["status", "summary", "updated_at"])
    return len(created), len(updated), skipped


def template_xlsx():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Foods"
    ws.append(["Food", "Portion_g", "Calories_kcal", "Portion", "Protein", "Carbs", "Fat", "Fibre", "Group", "Diet"])
    ws.append(["Paneer", 100, 265, "1 cube (25 g)", 18.3, 1.2, 20.8, 0, "Dairy", "Vegetarian"])
    ws.append(["Boiled egg", 50, 78, "1 egg (50 g)", 6.3, 0.6, 5.3, 0, "Egg", "Non-Vegetarian"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@transaction.atomic
def undo(job):
    """Toast Undo after an import: new foods go, updated foods get their old numbers back."""
    res = job.summary.get("result") or {}
    Food.objects.filter(id__in=res.get("created", [])).delete()
    for fid, prev in (res.get("previous") or {}).items():
        food = Food.objects.filter(id=fid).first()
        if food:
            for k, v in prev.items():
                setattr(food, k, Decimal(v) if v is not None else None)
            food.save()
    job.status = FoodImportJob.Status.FAILED
    job.summary = {**job.summary, "undone": True}
    job.save(update_fields=["status", "summary", "updated_at"])
