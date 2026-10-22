import io

from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from apps.nutrition.models import Food

HEAD = ["Category", "Food", "Portion/Quantity", "Portion_g", "Protein_g", "Carbs_g", "Fat_g", "Calories_kcal",
        "Fiber_g", "Sodium_mg", "Iron_mg", "Calcium_mg", "Potassium_mg", "Source", "Main_Category"]


def _xlsx(rows):
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile("foods.xlsx", buf.getvalue())


def test_template_has_every_column_of_vimals_sheet(admin):
    r = admin.get("/api/v1/foods/import/template/")
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    assert [c.value for c in ws[1]] == HEAD
    assert ws.max_row >= 3


def test_vimals_sheet_imports_minerals_group_diet_and_source(admin):
    f = _xlsx([HEAD,
               ["Legumes", "Test moong dal (cooked)", "1 cup (150 g)", 150, 9, 23, 3.5, 180, 7, 350, 2, 40, 350,
                "IFCT", "Vegetarian"],
               ["Chicken", "Test chicken tikka", "100 g", 100, 31, 0, 3.6, 165, 0, 74, 1, 15, 256, "USDA",
                "Non-Vegetarian"]])
    r = admin.post("/api/v1/foods/import/check/", {"file": f}, format="multipart")
    assert r.status_code == 200, r.content
    body = r.json()
    assert body["summary"]["new"] == 2 and body["summary"]["errors"] == 0
    r = admin.post(f"/api/v1/foods/import/{body['id']}/run/", {"duplicates": "skip"}, format="json")
    assert r.status_code == 200, r.content
    dal = Food.objects.get(name="Test moong dal (cooked)")
    assert (dal.group, dal.diet, dal.source_note) == ("protein", "veg", "IFCT")
    assert float(dal.sodium_mg) == 233.33 and float(dal.potassium_mg) == 233.33 and float(dal.fibre) == 4.67
    assert dal.default_unit.name == "1 cup" and float(dal.default_unit.grams) == 150
    tikka = Food.objects.get(name="Test chicken tikka")
    assert (tikka.group, tikka.diet, tikka.source_note) == ("protein", "nonveg", "USDA")
    assert tikka.default_unit is None  # "100 g" is not a unit


def test_headers_with_brackets_still_match(admin):
    f = _xlsx([["Food", "Portion (g)", "Calories (kcal)", "Sodium (mg)"], ["Test plain curd", 100, 60, 46]])
    r = admin.post("/api/v1/foods/import/check/", {"file": f}, format="multipart")
    assert r.status_code == 200, r.content
    assert r.json()["rows"][0]["per100"]["sodium_mg"] == 46
